from flask import Flask, render_template, request, send_file, redirect, url_for
from docx import Document
from docx2pdf import convert
import os
import uuid
import contextlib

app = Flask(__name__)
UPLOAD_FOLDER = "uploads"
RESULT_FOLDER = "results"

os.makedirs(UPLOAD_FOLDER, exist_ok=True)
os.makedirs(RESULT_FOLDER, exist_ok=True)

@app.route("/")
def index():
    return render_template("index.html")

@app.route("/generate", methods=["POST"])
def generate():
    name = request.form["name"]
    address = request.form["address"]
    phone = request.form["phone"]

    # Create .docx
    doc = Document()
    doc.add_heading("User Info", level=1)
    doc.add_paragraph(f"Name: {name}")
    doc.add_paragraph(f"Address: {address}")
    doc.add_paragraph(f"Phone: {phone}")

    # Unique filenames
    file_id = str(uuid.uuid4())
    docx_path = os.path.join(UPLOAD_FOLDER, f"{file_id}.docx")
    pdf_path = os.path.join(RESULT_FOLDER, f"{file_id}.pdf")

    doc.save(docx_path)

    # Convert to PDF (suppress tqdm error)
    with open(os.devnull, 'w') as fnull, contextlib.redirect_stdout(fnull), contextlib.redirect_stderr(fnull):
        convert(docx_path)

    # Move converted PDF to expected location
    generated_pdf = os.path.splitext(docx_path)[0] + ".pdf"
    if os.path.exists(generated_pdf):
        os.rename(generated_pdf, pdf_path)

    return send_file(pdf_path, as_attachment=True)

if __name__ == "__main__":
    app.run(debug=True)
